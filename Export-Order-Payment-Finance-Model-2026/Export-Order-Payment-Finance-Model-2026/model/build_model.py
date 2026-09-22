from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

NAVY="1B2A4A"; GOLD="C9A227"
wb=Workbook()
def hdr(ws,row,cols,text=None):
    for c in cols:
        cell=ws.cell(row=row,column=c)
        cell.font=Font(name="Arial",bold=True,color="FFFFFF",size=10)
        cell.fill=PatternFill("solid",fgColor=NAVY)
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
thin=Side(style="thin",color="BBBBBB"); BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
INPUT=Font(name="Arial",color="0000FF",size=10); F=Font(name="Arial",size=10); B=Font(name="Arial",bold=True,size=10)
YEL=PatternFill("solid",fgColor="FFFF00"); GOLDF=PatternFill("solid",fgColor="FFF2CC")
def title(ws,text,sub=None):
    ws["A1"]=text; ws["A1"].font=Font(name="Arial",bold=True,size=14,color=NAVY)
    ws["A2"]=sub or "Eswar Mahalingam · International Business Internship (Internship Studio) · Module 6 Project · Export Order Payment & Finance Model"
    ws["A2"].font=Font(name="Arial",italic=True,size=9,color="555555")

# ---------------- Assumptions ----------------
ws=wb.active; ws.title="Assumptions"
title(ws,"ASSUMPTIONS BLOCK — every other sheet references these cells")
ws["A4"]="Legend: blue = given input (from brief), yellow fill = editable scenario input, black = formula"; ws["A4"].font=Font(name="Arial",italic=True,size=9)
rows=[
("TRANSACTION FACTS",None,None,None),
("Export order value (FOB), USD",600000,"USD","Brief 3.1"),
("Production cost / packing credit base, USD",420000,"USD","Brief 3.1"),
("Pre-shipment (production) period, days",60,"days","Brief 3.1"),
("Credit period offered to buyer, days after shipment",90,"days","Brief 3.1"),
("Reporting exchange rate, INR per USD",83.5,"INR/USD","Brief 3.1 (fixed)"),
("Day-count convention, days per year",360,"days","Brief 3.1"),
("Advance payment price concession",0.05,"%","Brief 3.4"),
("BANK CHARGES",None,None,None),
("LC advising charge, USD (flat)",150,"USD","Brief 3.2"),
("LC confirmation charge, % per month on LC value",0.0015,"%/month","Brief 3.2"),
("Confirmation tenor, months (LC opened at order, valid to shipment = 60 days)",2,"months","Assumption: 60-day production = 2 months"),
("Negotiation / handling charge, % of bill value",0.002,"%","Brief 3.2"),
("SWIFT / courier, USD (flat)",120,"USD","Brief 3.2"),
("Documentary collection (D/A) handling, % of bill value",0.0015,"%","Brief 3.2"),
("Open account remittance handling, USD (flat)",150,"USD","Brief 3.2"),
("Inward remittance charge (advance), USD (flat)",100,"USD","Brief 3.2"),
("Discrepancy fee (LC), USD per set",90,"USD","Brief 3.2"),
("FINANCE RATES (per year)",None,None,None),
("Pre-shipment packing credit rate",0.09,"%","Brief 3.3"),
("Post-shipment credit / bill discounting (INR) rate",0.095,"%","Brief 3.3"),
("Interest subvention benefit (reduction)",0.03,"%","Brief 3.3"),
("Foreign currency usance bill discount rate",0.065,"%","Brief 3.3 (already net)"),
("Exporter opportunity cost of capital (PV rate)",0.12,"%","Brief 3.3"),
("Subvented packing credit rate",  "=B25-B27","%","Formula: base − subvention"),
("Subvented post-shipment credit rate","=B26-B27","%","Formula: base − subvention"),
("DEFAULT PROBABILITIES",None,None,None),
("Advance payment (100%)",0.0,"%","Brief 3.4"),
("Confirmed LC at sight",0.002,"%","Brief 3.4"),
("Usance LC 90 days (unconfirmed)",0.01,"%","Brief 3.4"),
("Documentary collection D/A 90 days",0.04,"%","Brief 3.4"),
("Open account 90 days",0.08,"%","Brief 3.4"),
("TASK E SCENARIO INPUTS",None,None,None),
("Number of discrepant document sets presented",1,"sets","Task E scenario"),
("Default probability if LC is discrepant (bank undertaking lost → buyer credit risk, open-account level)",0.08,"%","Scenario (editable)"),
("Extra days delay while buyer waiver is sought",10,"days","Scenario (editable)"),
]
ws["A5"]="Item"; ws["B5"]="Value"; ws["C5"]="Unit"; ws["D5"]="Source / note"; hdr(ws,5,[1,2,3,4])
r=6
for it,val,unit,src in rows:
    ws.cell(r,1,it); 
    if val is None:
        ws.cell(r,1).font=Font(name="Arial",bold=True,color=NAVY,size=10); ws.cell(r,1).fill=GOLDF
    else:
        c=ws.cell(r,2,val); c.font=F if str(val).startswith("=") else INPUT
        if "Scenario" in (src or "") or "Assumption" in (src or ""): c.fill=YEL
        ws.cell(r,3,unit).font=F; ws.cell(r,4,src).font=Font(name="Arial",size=9,color="555555")
        if unit and "%" in unit: c.number_format="0.00%"
        elif unit=="USD": c.number_format='$#,##0'
        elif unit=="INR/USD": c.number_format="0.00"
    for cc in range(1,5): ws.cell(r,cc).border=BORDER
    r+=1
# verify row numbers of key cells
names={}
for rr in range(6,r):
    names[ws.cell(rr,1).value]=rr
A=lambda label: f"Assumptions!$B${names[label]}"
ws.column_dimensions["A"].width=78; ws.column_dimensions["B"].width=14; ws.column_dimensions["C"].width=10; ws.column_dimensions["D"].width=42
ws.freeze_panes="A6"

ORDER=A("Export order value (FOB), USD"); PCBASE=A("Production cost / packing credit base, USD"); PDAYS=A("Pre-shipment (production) period, days")
CDAYS=A("Credit period offered to buyer, days after shipment"); FX=A("Reporting exchange rate, INR per USD"); DC=A("Day-count convention, days per year")
CONC=A("Advance payment price concession"); ADV=A("LC advising charge, USD (flat)"); CONFP=A("LC confirmation charge, % per month on LC value")
CONFM=A("Confirmation tenor, months (LC opened at order, valid to shipment = 60 days)"); NEG=A("Negotiation / handling charge, % of bill value")
SWIFT=A("SWIFT / courier, USD (flat)"); DAH=A("Documentary collection (D/A) handling, % of bill value"); OAH=A("Open account remittance handling, USD (flat)")
INW=A("Inward remittance charge (advance), USD (flat)"); DISCF=A("Discrepancy fee (LC), USD per set")
PCR=A("Subvented packing credit rate"); PSR=A("Subvented post-shipment credit rate"); FCR=A("Foreign currency usance bill discount rate"); COC=A("Exporter opportunity cost of capital (PV rate)")
PD=[A("Advance payment (100%)"),A("Confirmed LC at sight"),A("Usance LC 90 days (unconfirmed)"),A("Documentary collection D/A 90 days"),A("Open account 90 days")]
NSETS=A("Number of discrepant document sets presented"); PDDISC=A("Default probability if LC is discrepant (bank undertaking lost → buyer credit risk, open-account level)"); DELAY=A("Extra days delay while buyer waiver is sought")

# ---------------- Task A ----------------
ws=wb.create_sheet("TaskA_PackingCredit")
title(ws,"TASK A — Pre-shipment financing cost (packing credit)")
ws["A4"]="Step"; ws["B4"]="Item"; ws["C4"]="Value"; ws["D4"]="Formula (as typed)"; hdr(ws,4,[1,2,3,4])
ta=[
("1","Packing credit principal, USD",f"={PCBASE}","=Assumptions!B8",'$#,##0'),
("2","Base packing credit rate",f"={A('Pre-shipment packing credit rate')}","=Assumptions!B25","0.00%"),
("3","Interest subvention",f"={A('Interest subvention benefit (reduction)')}","=Assumptions!B27","0.00%"),
("4","Subvented rate (base − subvention)","=C6-C7","=C6-C7","0.00%"),
("5","Production period, days",f"={PDAYS}","=Assumptions!B9","0"),
("6","Day count",f"={DC}","=Assumptions!B12","0"),
("7","Packing credit interest, USD = Principal × subvented rate × days/360","=C5*C8*C9/C10","=C5*C8*C9/C10",'$#,##0.00'),
("8","Packing credit interest, INR (× 83.50)",f"=C11*{FX}","=C11*Assumptions!B11",'"₹"#,##0'),
("9","Interest without subvention (for comparison), USD","=C5*C6*C9/C10","=C5*C6*C9/C10",'$#,##0.00'),
("10","Subvention saving, USD","=C13-C11","=C13-C11",'$#,##0.00'),
]
for i,(s,it,f,ft,nf) in enumerate(ta):
    rr=5+i; ws.cell(rr,1,s); ws.cell(rr,2,it); ws.cell(rr,3,f).number_format=nf; fc=ws.cell(rr,4,ft); fc.data_type="s"; fc.font=Font(name="Courier New",size=9)
    for cc in range(1,5): ws.cell(rr,cc).border=BORDER; 
    ws.cell(rr,1).font=F; ws.cell(rr,2).font=F; ws.cell(rr,3).font=B if s=="7" else F
ws["A16"]="Note: this cost applies to every payment term except advance payment (cash arrives before production, so no packing credit is drawn)."; ws["A16"].font=Font(name="Arial",italic=True,size=9)
ws.column_dimensions["A"].width=6; ws.column_dimensions["B"].width=62; ws.column_dimensions["C"].width=16; ws.column_dimensions["D"].width=30

# ---------------- Task B/C: Model ----------------
ws=wb.create_sheet("TaskB_C_Model")
title(ws,"TASK B & C — Cost build-up, net realisation, present value and all-in cost for the five payment terms")
terms=["Advance payment (100%)","Confirmed LC at sight","Usance LC 90 days (unconfirmed)","Documentary collection D/A 90 days","Open account 90 days"]
ws["A4"]="Row"; hdr(ws,4,[1]+list(range(2,7)))
for j,t in enumerate(terms): ws.cell(4,2+j,t)
ws.row_dimensions[4].height=32
cols=["B","C","D","E","F"]
lines=[]  # (label, [formulas per col], numfmt, bold)
M='$#,##0'
lines.append(("Special condition",["5% price concession demanded by buyer","Confirmation charge applies","Bill discounted in foreign currency @ 6.5%","No bank undertaking; post-shipment credit","No documents held; post-shipment credit"],None,False))
lines.append(("Payment timing (day)",["0",f"={PDAYS}",f"={PDAYS}",f"={PDAYS}",f"={PDAYS}"],"0",False))
lines.append(("Cash timing basis",["Advance received before production","Sight payment at shipment (day 60)","Discounted proceeds at shipment (day 60)","Credit-financed proceeds at shipment (day 60)","Credit-financed proceeds at shipment (day 60)"],None,False))
lines.append(("GROSS AMOUNT RECEIVED",[f"={ORDER}*(1-{CONC})",f"={ORDER}",f"={ORDER}",f"={ORDER}",f"={ORDER}"],M,True))
lines.append(("Bank charge: LC advising",["0",f"={ADV}",f"={ADV}","0","0"],M,False))
lines.append(("Bank charge: LC confirmation (0.15%/month × months × LC value)",["0",f"={CONFP}*{CONFM}*{ORDER}","0","0","0"],M,False))
lines.append(("Bank charge: negotiation / handling (0.20% × bill)",["0",f"={NEG}*{ORDER}",f"={NEG}*{ORDER}","0","0"],M,False))
lines.append(("Bank charge: D/A collection handling (0.15% × bill)",["0","0","0",f"={DAH}*{ORDER}","0"],M,False))
lines.append(("Bank charge: open account remittance handling",["0","0","0","0",f"={OAH}"],M,False))
lines.append(("Bank charge: inward remittance (advance)",[f"={INW}","0","0","0","0"],M,False))
lines.append(("Bank charge: SWIFT / courier",["0",f"={SWIFT}",f"={SWIFT}",f"={SWIFT}",f"={SWIFT}"],M,False))
lines.append(("TOTAL BANK CHARGES",[f"=SUM({c}9:{c}15)" for c in cols],M,True))
lines.append(("Financing / discounting cost (usance LC: bill × 6.5% × 90/360; D/A & OA: bill × subvented 6.5% × 90/360)",["0","0",f"={ORDER}*{FCR}*{CDAYS}/{DC}",f"={ORDER}*{PSR}*{CDAYS}/{DC}",f"={ORDER}*{PSR}*{CDAYS}/{DC}"],M,False))
lines.append(("Packing credit cost (Task A)",["0"]+["=TaskA_PackingCredit!$C$11"]*4,M,False))
lines.append(("Default probability",[f"={p}" for p in PD],"0.0%",False))
lines.append(("Expected loss = probability × order value",[f"={c}19*{ORDER}" for c in cols],M,False))
lines.append(("NET REALISATION (nominal) = gross − charges − financing − packing credit − expected loss",[f"={c}8-{c}16-{c}17-{c}18-{c}20" for c in cols],M,True))
lines.append(("Net realisation, INR (× 83.50)",[f"={c}21*{FX}" for c in cols],'"₹"#,##0',False))
lines.append(("Days from day 0 to cash",[f"={c}6" for c in cols],"0",False))
lines.append(("PV factor (simple) = 1 / (1 + 12% × days/360)",[f"=1/(1+{COC}*{c}23/{DC})" for c in cols],"0.000000",False))
lines.append(("PRESENT VALUE OF NET REALISATION, USD",[f"={c}21*{c}24" for c in cols],M,True))
lines.append(("Present value, INR",[f"={c}25*{FX}" for c in cols],'"₹"#,##0',False))
lines.append(("ALL-IN COST % = (order value − PV) / order value",[f"=({ORDER}-{c}25)/{ORDER}" for c in cols],"0.00%",True))
lines.append(("All-in cost, USD (order value − PV)",[f"={ORDER}-{c}25" for c in cols],M,False))
lines.append(("Rank by PV of net realisation (1 = best)",[f"=RANK({c}25,$B$25:$F$25,0)" for c in cols],"0",True))
lines.append(("Rank by default risk (1 = safest)",[f"=RANK({c}19,$B$19:$F$19,1)" for c in cols],"0",True))
lines.append(("Cost of risk share: expected loss as % of all-in cost",[f"=IF({c}28=0,0,{c}20/{c}28)" for c in cols],"0.0%",False))
for i,(lab,fs,nf,bold) in enumerate(lines):
    rr=5+i; ws.cell(rr,1,lab).font=B if bold else F
    if bold: ws.cell(rr,1).fill=GOLDF
    for j,f in enumerate(fs):
        c=ws.cell(rr,2+j)
        c.value=(float(f) if f.replace('.','',1).isdigit() else f)
        if nf: c.number_format=nf
        c.font=B if bold else F
        if bold: c.fill=GOLDF
        c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="right" if nf else "left")
    for cc in range(1,7): ws.cell(rr,cc).border=BORDER
ws.column_dimensions["A"].width=64
for c in cols: ws.column_dimensions[c].width=24
ws.row_dimensions[5].height=30; ws.row_dimensions[7].height=30; ws.row_dimensions[17].height=42; ws.row_dimensions[21].height=30
ws.freeze_panes="B5"
ws["A37"]="Checks"; ws["A37"].font=B
ws["A38"]="Advance premium check: does the 5% concession (USD) exceed total sight-LC frictions (charges + packing credit + expected loss)?"; ws["B38"]=f"={ORDER}*{CONC}"; ws["B38"].number_format=M; ws["C38"]="=C16+C18+C20"; ws["C38"].number_format=M; ws["D38"]='=IF(B38>C38,"YES — concession costs more than sight LC","NO")'
ws["A39"]="Best term by PV (formula)"; ws["B39"]="=INDEX($B$4:$F$4,MATCH(1,$B$29:$F$29,0))"
ws["A40"]="Safest term by default risk (formula)"; ws["B40"]="=INDEX($B$4:$F$4,MATCH(1,$B$30:$F$30,0))"
ws["A41"]="PV gap: best term minus second-best, USD"; ws["B41"]="=LARGE($B$25:$F$25,1)-LARGE($B$25:$F$25,2)"; ws["B41"].number_format=M
for rr in range(38,42):
    for cc in range(1,5): ws.cell(rr,cc).border=BORDER; ws.cell(rr,cc).font=F

# ---------------- Task D summary + chart ----------------
ws=wb.create_sheet("TaskD_Summary")
title(ws,"TASK D — Ranking, comparison table and chart")
hd=["Payment term","Default probability","Expected loss (USD)","Net realisation nominal (USD)","PV of net realisation (USD)","PV (INR)","All-in cost %","Rank by PV","Rank by risk","Verdict"]
for j,h in enumerate(hd): ws.cell(4,1+j,h)
hdr(ws,4,list(range(1,11))); ws.row_dimensions[4].height=32
for i,c in enumerate(cols):
    rr=5+i
    ws.cell(rr,1,f"=TaskB_C_Model!{c}4")
    ws.cell(rr,2,f"=TaskB_C_Model!{c}19").number_format="0.0%"
    ws.cell(rr,3,f"=TaskB_C_Model!{c}20").number_format=M
    ws.cell(rr,4,f"=TaskB_C_Model!{c}21").number_format=M
    ws.cell(rr,5,f"=TaskB_C_Model!{c}25").number_format=M
    ws.cell(rr,6,f"=TaskB_C_Model!{c}26").number_format='"₹"#,##0'
    ws.cell(rr,7,f"=TaskB_C_Model!{c}27").number_format="0.00%"
    ws.cell(rr,8,f"=TaskB_C_Model!{c}29")
    ws.cell(rr,9,f"=TaskB_C_Model!{c}30")
    ws.cell(rr,10,f'=IF(H{rr}=1,"RECOMMENDED",IF(H{rr}=2,"Runner-up",IF(B{rr}>=0.04,"Reject — risk too high","Acceptable fallback")))')
    for cc in range(1,11): ws.cell(rr,cc).border=BORDER; ws.cell(rr,cc).font=Font(name="Arial",size=10,color="008000") if cc<10 else B
ws["A11"]="Recommended structure:"; ws["A11"].font=B; ws["B11"]="=TaskB_C_Model!B39"; ws["B11"].font=Font(name="Arial",bold=True,color=NAVY,size=11)
ws["A12"]="Advance PV shortfall vs recommended, USD:"; ws["A12"].font=B; ws["B12"]="=E6-E5"; ws["B12"].number_format=M
ws["A13"]="Open account PV shortfall vs recommended, USD:"; ws["A13"].font=B; ws["B13"]="=E6-E9"; ws["B13"].number_format=M
ws["A14"]="Is advance payment really best once the concession is counted?"; ws["A14"].font=B; ws["B14"]='=IF(E5>=MAX(E5:E9),"Yes","No — the 5% concession (USD 30,000) costs more than every friction on the sight LC combined")'
for c in "ABCDEFGHIJ": ws.column_dimensions[c].width=20
ws.column_dimensions["A"].width=40; ws.column_dimensions["J"].width=24
ch=BarChart(); ch.type="col"; ch.title="PV of net realisation by payment term (USD)"; ch.y_axis.title="USD"; ch.height=9; ch.width=20
ch.add_data(Reference(ws,min_col=5,min_row=4,max_row=9),titles_from_data=True); ch.set_categories(Reference(ws,min_col=1,min_row=5,max_row=9)); ch.legend=None
ws.add_chart(ch,"A17")
ch2=BarChart(); ch2.type="col"; ch2.title="All-in cost % vs default probability"; ch2.height=9; ch2.width=20
ch2.add_data(Reference(ws,min_col=7,min_row=4,max_row=9),titles_from_data=True); ch2.add_data(Reference(ws,min_col=2,min_row=4,max_row=9),titles_from_data=True); ch2.set_categories(Reference(ws,min_col=1,min_row=5,max_row=9))
ws.add_chart(ch2,"F17")

# ---------------- Task E ----------------
ws=wb.create_sheet("TaskE_UCP600")
title(ws,"TASK E — UCP 600 documentary check on the confirmed sight LC presentation and its financial impact")
hd=["#","Document / item presented","LC requirement","Presented as","Discrepancy?","UCP 600 basis","Why it fails"]
for j,h in enumerate(hd): ws.cell(4,1+j,h)
hdr(ws,4,list(range(1,8))); ws.row_dimensions[4].height=30
te=[
(1,"Commercial invoice","Signed invoice for the LC amount (USD 600,000)","Shows USD 605,000","YES","Art. 18(b) — a nominated/confirming bank may refuse an invoice issued for an amount in excess of the credit; Art. 14(d) data must not conflict","Invoice exceeds the LC amount by USD 5,000; no ±tolerance wording (\"about\"/\"approximately\") in the LC (Art. 30)."),
(2,"Ocean bill of lading","Full set of CLEAN SHIPPED ON BOARD B/L","Marked \"received for shipment\"","YES","Art. 20(a)(ii) — B/L must indicate goods shipped on board a named vessel by pre-printed wording or an on-board notation with date","\"Received for shipment\" only proves receipt by the carrier, not loading; no on-board date, so shipment cannot be evidenced."),
(3,"Marine insurance policy","Cover for 110% of CIF value","Covers 100% of value","YES","Art. 28(f)(ii) — minimum cover is 110% of CIF/CIP value when the credit so requires (and even by default)","Under-insured by 10 percentage points; the bank/buyer are not protected for the full landed value plus expected profit."),
(4,"Certificate of origin","Consistent with other documents","Consistent","NO","Art. 14(d)/(e) — data need not be identical but must not conflict","Complies — no action."),
(5,"Shipment date","On or before the LC's latest shipment date","Later than the latest shipment date","YES","Art. 14(c) and the LC's latest-shipment field (MT700 44C); late shipment is a classic discrepancy","Documents dated after the latest shipment date cannot be honoured; also risks the Art. 14(c) 21-day presentation window and LC expiry."),
(6,"Packing list & sight bill of exchange","Required","Presented (no issue stated)","NO","Art. 14","Complies — no action."),
]
for i,row in enumerate(te):
    rr=5+i
    for j,v in enumerate(row):
        c=ws.cell(rr,1+j,v); c.font=F; c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=BORDER
        if j==4: c.font=Font(name="Arial",bold=True,color=("C00000" if v=="YES" else "008000"))
    ws.row_dimensions[rr].height=58
ws["A12"]="Number of discrepancies found"; ws["B12"]='=COUNTIF(E5:E10,"YES")'; ws["A12"].font=B
ws["A13"]="Discrepancy fee, USD (USD 90 × sets)"; ws["B13"]=f"={DISCF}*{NSETS}"; ws["B13"].number_format=M; ws["A13"].font=B
ws["A15"]="FINANCIAL IMPACT ON THE SIGHT LC OPTION"; ws["A15"].font=Font(name="Arial",bold=True,color=NAVY,size=11); ws["A15"].fill=GOLDF
hd2=["Line","Compliant presentation","Discrepant presentation","Change","Explanation"]
for j,h in enumerate(hd2): ws.cell(16,1+j,h)
hdr(ws,16,[1,2,3,4,5])
imp=[
("Default probability","=TaskB_C_Model!C19",f"={PDDISC}","=C17-B17","Bank undertaking is lost (Art. 16 refusal) — payment now depends on the buyer's waiver, i.e. the buyer's own credit risk (open-account level).","0.0%"),
("Expected loss, USD",f"=B17*{ORDER}",f"=C17*{ORDER}","=C18-B18","Probability × exposure.",M),
("Discrepancy fee, USD","0","=B13","=C19-B19","USD 90 per set presented.",M),
("Days to cash",f"={PDAYS}",f"={PDAYS}+{DELAY}","=C20-B20","Waiver / re-presentation delay (scenario input).","0"),
("Net realisation (nominal), USD","=TaskB_C_Model!C21","=B21-D18-D19","=C21-B21","Extra expected loss and fee reduce the nominal figure.",M),
("PV factor","=TaskB_C_Model!C24",f"=1/(1+{COC}*C20/{DC})","=C22-B22","Longer wait lowers the PV factor.","0.000000"),
("PV of net realisation, USD","=B21*B22","=C21*C22","=C23-B23","Combined effect.",M),
("All-in cost %",f"=({ORDER}-B23)/{ORDER}",f"=({ORDER}-C23)/{ORDER}","=C24-B24","Sight LC drifts from best option to open-account-like economics.","0.00%"),
("Rank among the five terms (PV)","=TaskB_C_Model!C29","=1+COUNTIF(TaskB_C_Model!$B$25:$F$25,\">\"&C23)-IF(C23<TaskB_C_Model!C25,1,0)","=C25-B25","Recomputed rank of the discrepant LC against the other four terms.","0"),
]
for i,(lab,b,c_,d,e,nf) in enumerate(imp):
    rr=17+i; ws.cell(rr,1,lab).font=B; ws.cell(rr,2,b if b!="0" else 0).number_format=nf; ws.cell(rr,3,c_).number_format=nf; ws.cell(rr,4,d).number_format=nf; ws.cell(rr,5,e).alignment=Alignment(wrap_text=True,vertical="top")
    for cc in range(1,6): ws.cell(rr,cc).border=BORDER
    for cc in range(2,6): ws.cell(rr,cc).font=F
    ws.row_dimensions[rr].height=30
ws["A27"]="Cure actions: (1) re-issue invoice at USD 600,000; (2) obtain on-board notation with date on the B/L; (3) endorse insurance to 110% CIF; (5) request LC amendment extending latest shipment date and expiry BEFORE presenting — then the presentation is compliant and the confirming bank's undertaking is restored."
ws["A27"].alignment=Alignment(wrap_text=True); ws.merge_cells("A27:G27"); ws.row_dimensions[27].height=40; ws["A27"].font=F
ws.column_dimensions["A"].width=30; ws.column_dimensions["B"].width=34; ws.column_dimensions["C"].width=30; ws.column_dimensions["D"].width=14; ws.column_dimensions["E"].width=48; ws.column_dimensions["F"].width=44; ws.column_dimensions["G"].width=48

# ---------------- Sensitivity ----------------
ws=wb.create_sheet("Sensitivity")
title(ws,"SENSITIVITY — how the ranking moves with cost of capital and open-account default risk (all live formulas)")
ws["A4"]="PV of net realisation (USD) at different exporter costs of capital"; ws["A4"].font=B
ws["A5"]="Payment term"; hdr(ws,5,[1,2,3,4,5,6])
cocs=[0.08,0.10,0.12,0.14,0.16]
for j,v in enumerate(cocs): c=ws.cell(5,2+j,v); c.number_format="0%"; c.font=Font(name="Arial",bold=True,color="FFFFFF"); 
for i,c in enumerate(cols):
    rr=6+i; ws.cell(rr,1,f"=TaskB_C_Model!{c}4")
    for j in range(5):
        col=get_column_letter(2+j)
        ws.cell(rr,2+j,f"=TaskB_C_Model!{c}21/(1+{col}$5*TaskB_C_Model!{c}23/{DC})").number_format=M
    for cc in range(1,7): ws.cell(rr,cc).border=BORDER; ws.cell(rr,cc).font=F
ws["A12"]="Best term at each rate"; ws["A12"].font=B
for j in range(5):
    col=get_column_letter(2+j); ws.cell(12,2+j,f"=INDEX($A$6:$A$10,MATCH(MAX({col}6:{col}10),{col}6:{col}10,0))").font=Font(name="Arial",bold=True,color=NAVY); ws.cell(12,2+j).alignment=Alignment(wrap_text=True)
ws.row_dimensions[12].height=30
ws["A14"]="Break-even default probability at which open account would equal the sight LC on PV"; ws["A14"].font=B
ws["A15"]="OA net realisation before expected loss, USD"; ws["B15"]="=TaskB_C_Model!F8-TaskB_C_Model!F16-TaskB_C_Model!F17-TaskB_C_Model!F18"; ws["B15"].number_format=M
ws["A16"]="Required OA net realisation to match sight-LC PV, USD"; ws["B16"]="=TaskB_C_Model!C25/TaskB_C_Model!F24"; ws["B16"].number_format=M
ws["A17"]="Break-even default probability"; ws["B17"]=f"=(B15-B16)/{ORDER}"; ws["B17"].number_format="0.00%"
ws["A18"]="Interpretation"; ws["B18"]='=IF(B17<0,"Open account can never beat the sight LC even at zero default risk","Open account only wins if the buyer default probability is below the break-even shown")'
ws["A20"]="Advance concession break-even: concession % at which advance payment equals sight-LC PV"; ws["A20"].font=B
ws["B20"]=f"=1-(TaskB_C_Model!C25+{INW})/{ORDER}"; ws["B20"].number_format="0.00%"
ws["A21"]="Interpretation"; ws["B21"]='="Advance beats the sight LC only if the buyer\'s concession is below "&TEXT(B20,"0.00%")&" (buyer demands 5.00%)"'
for rr in range(15,22):
    for cc in range(1,3): ws.cell(rr,cc).border=BORDER; ws.cell(rr,cc).font=F
ws.column_dimensions["A"].width=70
for c in "BCDEF": ws.column_dimensions[c].width=22

for w in wb.worksheets:
    w.page_setup.orientation="landscape"; w.page_setup.fitToWidth=1; w.page_setup.fitToHeight=0; w.sheet_properties.pageSetUpPr.fitToPage=True
    w.oddFooter.center.text="&A — Eswar Mahalingam · IB Module 6 · Page &P"
wb.save("Export_Payment_Finance_Model_Eswar.xlsx")
print("saved", names)
